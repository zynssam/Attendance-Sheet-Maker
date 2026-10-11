#!/usr/bin/env python3
"""
build.py — Attendance Sheet Maker build script.

Reads the face members and 3 attendance templates from data/,
embeds them into template.html, and writes attendance_app.html.

Re-run this script whenever the /data files change.
"""

import base64
import json
import re
from pathlib import Path

import openpyxl
from jinja2 import Environment, FileSystemLoader

# ── Paths ────────────────────────────────────────────────────────────────────

DATA_DIR = Path("data")
TEMPLATE_FILE = "template.html"
OUTPUT_FILE = Path("attendance_app.html")

# ── Constants ─────────────────────────────────────────────────────────────────

COORDINATORS = {
    1: "Prof. Mala B M",
    2: "Prof. Ravi B K",
    3: "Prof. Manasa K N",
}

ATTENDANCE_FILES = {
    1: DATA_DIR / "Attendance_21-9_Group1_Badging.xlsx",
    2: DATA_DIR / "Attendance_21-9_Group2_Badging.xlsx",
    3: DATA_DIR / "Attendance_21-9_Group3_Badging.xlsx",
}

FACE_MEMBERS_FILE = DATA_DIR / "Facemembers_2026-27_Grouped.xlsx"

# Manual data corrections: USN -> {field: corrected_value}
CORRECTIONS = {
    "R25EH122": {"branch": "AIDS C"},  # Shivanshi — was incorrectly "CSE"
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def extract_section(branch):
    """
    Extract the class section letter (A-F) from a branch/department string.
    Takes the LAST standalone A-F letter (not embedded in a longer word).

    Examples:
      'AIDS C'      -> 'C'
      'AIML-E'      -> 'E'
      'AI&DS-A'     -> 'A'
      'CSE core E'  -> 'E'
      'IOT'         -> ''
      'CSE'         -> ''
    """
    if not branch:
        return ""
    matches = re.findall(r"(?<![A-Za-z])([A-F])(?![A-Za-z])", branch.upper())
    return matches[-1] if matches else ""


def parse_sheet_group_sem(sheet_name):
    """
    Parse group number and semester from a face-members sheet name.
    Format: 'Group N - Mth Sem'
    Returns (group_num, semester_num) or (None, None).
    """
    m = re.match(r"Group\s+(\d+)\s+-\s+(\d+)", sheet_name)
    if not m:
        return None, None
    return int(m.group(1)), int(m.group(2))


# ── Data builders ─────────────────────────────────────────────────────────────

def build_face_members():
    """
    Read Facemembers_2026-27_Grouped.xlsx.
    Returns a dict keyed by uppercase USN.
    """
    wb = openpyxl.load_workbook(FACE_MEMBERS_FILE)
    members = {}

    for sheet_name in wb.sheetnames:
        group, semester = parse_sheet_group_sem(sheet_name)
        if group is None:
            continue

        ws = wb[sheet_name]
        # Columns: Sl.No | Section | Name | USN | Committee | Branch | Role
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 4 or not row[3]:
                continue

            usn = str(row[3]).strip().upper()
            name = str(row[2]).strip() if row[2] else ""
            branch = str(row[5]).strip() if len(row) > 5 and row[5] else ""

            # Apply manual corrections
            if usn in CORRECTIONS:
                correction = CORRECTIONS[usn]
                if "branch" in correction:
                    branch = correction["branch"]

            section = extract_section(branch)

            members[usn] = {
                "name": name,
                "group": group,
                "semester": semester,
                "branch": branch,
                "section": section,
            }

    return members


def build_section_fallback(face_members):
    """
    Build a USN -> section fallback map from the old attendance template sheets.
    Provides a section for students whose branch has no section letter.
    """
    fallback = {}

    for _g, fpath in ATTENDANCE_FILES.items():
        wb = openpyxl.load_workbook(fpath)
        for sheet_name in wb.sheetnames:
            if "Link" in sheet_name:
                continue
            ws = wb[sheet_name]
            # Data starts at row 3; columns: SL.No | SRN | School | Semester | Section | ...
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not row or row[0] is None:
                    continue
                usn = str(row[1]).strip().upper() if row[1] else ""
                section = str(row[4]).strip() if len(row) > 4 and row[4] else ""
                if usn and section:
                    fallback.setdefault(usn, section)

    return fallback


def b64_file(path):
    """Read a binary file and return its base64-encoded ASCII string."""
    return base64.b64encode(path.read_bytes()).decode("ascii")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("Reading face members ...")
    face_members = build_face_members()
    print(f"  -> {len(face_members)} members loaded")

    shivanshi = face_members.get("R25EH122")
    if shivanshi:
        print(f"  ok Shivanshi R25EH122: branch='{shivanshi['branch']}', section='{shivanshi['section']}'")

    print("Building section fallback ...")
    section_fallback = build_section_fallback(face_members)
    print(f"  -> {len(section_fallback)} USN->section entries")

    print("Base64-encoding attendance templates ...")
    b64 = {}
    for g, fpath in ATTENDANCE_FILES.items():
        b64[g] = b64_file(fpath)
        print(f"  -> Group {g}: {fpath.name}  ({len(b64[g]):,} chars base64)")

    print("Rendering HTML template ...")
    env = Environment(loader=FileSystemLoader("."), autoescape=False)
    template = env.get_template(TEMPLATE_FILE)

    html = template.render(
        face_members_json=json.dumps(face_members, ensure_ascii=False),
        section_fallback_json=json.dumps(section_fallback, ensure_ascii=False),
        b64_group1=b64[1],
        b64_group2=b64[2],
        b64_group3=b64[3],
        coordinators_json=json.dumps(COORDINATORS),
    )

    OUTPUT_FILE.write_text(html, encoding="utf-8")
    size_kb = OUTPUT_FILE.stat().st_size // 1024
    print(f"\nBuilt: {OUTPUT_FILE}  ({size_kb} KB)")
    print("  Open attendance_app.html in any modern browser.")


if __name__ == "__main__":
    main()
